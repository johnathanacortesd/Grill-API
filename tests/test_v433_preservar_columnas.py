"""v4.33 — Pestaña "Columnas personalizadas": conservar el xlsx como está.

- `construir_ai_config_custom` marca preservar_columnas=True.
- `_columnas_preservadas`: columnas originales en su orden + columnas nuevas
  del análisis al final (nunca el esquema fijo BASE_OUTPUT_COLUMNS).
- End-to-end: xlsx sintético estilo Colpensiones (primera columna ID,
  columnas LINK/WEB con la palabra "Link" hipervinculada) → el resultado
  conserva ID, los hipervínculos clicables y el orden de columnas.
"""
import io
import unittest

from openpyxl import Workbook, load_workbook

from pipeline import (
    _columnas_preservadas,
    construir_ai_config_custom,
    process_dossier,
    renombrar_columnas_xlsx,
)


def _cfg_custom(**kw):
    base = dict(brand="Colpensiones", alias_txt="", voceros_txt="",
                criterio="General", incluir_tema=False,
                incluir_prominencia=True, enable_ai=False, api_key=None,
                typesafe_api_key=None, historial_dir="/tmp",
                tone_pkl_bytes=None, theme_pkl_bytes=None)
    base.update(kw)
    return construir_ai_config_custom(**base)


def _dossier_estilo_colpensiones():
    """XLSX sintético: ID (primera col), TÍTULO, RESUMEN, LINK/WEB con 'Link'
    hipervinculado — como la base real de Colpensiones."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    headers = ["ID", "FECHA", "TÍTULO", "RESUMEN", "LINK", "WEB"]
    ws.append(headers)
    filas = [
        (61142010, "2026-09-01", "Gobierno avanza en ruta pensional",
         "El gobierno anunció la ruta técnica para las mesadas.",
         "https://ejemplo.co/nota1", "https://ejemplo.co/web1"),
        (61142304, "2026-09-02", "Ministra habla de pensiones",
         "La ministra del Trabajo explicó el faltante de billones.",
         "https://ejemplo.co/nota2", "https://ejemplo.co/web2"),
    ]
    for f in filas:
        ws.append(list(f))
    for r in range(2, 4):
        ws.cell(row=r, column=5).hyperlink = ws.cell(row=r, column=5).value
        ws.cell(row=r, column=5).value = "Link"
        ws.cell(row=r, column=6).hyperlink = ws.cell(row=r, column=6).value
        ws.cell(row=r, column=6).value = "Link"
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def _celdas_salida(xlsx_bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    data = list(ws.iter_rows(min_row=2, values_only=False))
    return headers, ws, data


class TestPreservarColumnasV433(unittest.TestCase):
    def test_config_custom_marca_preservar(self):
        cfg = _cfg_custom()
        self.assertTrue(cfg["preservar_columnas"])

    def test_columnas_preservadas_orden_y_agregados(self):
        cols = _columnas_preservadas(
            ["ID", "FECHA", "TÍTULO", "RESUMEN", "LINK", "WEB"],
            include_ai=True, include_tema=True, include_subtema=True,
            include_prominencia=True)
        self.assertEqual(cols[:6], ["ID", "FECHA", "TÍTULO", "RESUMEN", "LINK", "WEB"])
        self.assertEqual(cols[6:], ["Tono_IA", "Tema_IA", "Subtema_IA",
                                   "Contexto analizado", "Prominencia"])

    def test_columnas_preservadas_sin_ia(self):
        cols = _columnas_preservadas(["ID", "LINK"], include_ai=False,
                                    include_prominencia=True)
        self.assertEqual(cols, ["ID", "LINK", "Prominencia"])
        self.assertNotIn("Contexto analizado", cols)

    def test_columnas_preservadas_sin_subtema(self):
        cols = _columnas_preservadas(["ID"], include_ai=True,
                                    include_subtema=False)
        self.assertIn("Tono_IA", cols)
        self.assertIn("Tema_IA", cols)
        self.assertNotIn("Subtema_IA", cols)

    def test_columnas_preservadas_sin_duplicar(self):
        cols = _columnas_preservadas(["ID", "Tono_IA"], include_ai=True)
        self.assertEqual(cols.count("Tono_IA"), 1)
        self.assertEqual(cols[0], "ID")

    def test_e2e_conserva_id_e_hipervinculos(self):
        blob = _dossier_estilo_colpensiones()
        renombrado = renombrar_columnas_xlsx(blob, "TÍTULO", "RESUMEN")
        res = process_dossier(io.BytesIO(renombrado), {}, {},
                              ai_config=_cfg_custom())
        headers, ws, data = _celdas_salida(res["output_data"])
        # Columnas originales en su orden + Prominencia al final.
        self.assertEqual(headers, ["ID", "FECHA", "Título",
                                  "Resumen - Aclaracion", "LINK", "WEB",
                                  "Prominencia"])
        # La data de la primera columna se conserva.
        ids = sorted(str(r[0].value) for r in data)
        self.assertEqual(ids, ["61142010", "61142304"])
        # Los hipervínculos de la palabra "Link" siguen clicables.
        for r in data:
            for cidx, esperado in ((4, "nota"), (5, "web")):
                celda = r[cidx]
                self.assertEqual(celda.value, "Link")
                self.assertIsNotNone(celda.hyperlink,
                                     f"fila {r[0].value} col {headers[cidx]} sin hipervínculo")
                self.assertIn(esperado, celda.hyperlink.target)
        ws._parent.close()

    def test_e2e_sin_preservar_usa_esquema_base(self):
        # Sin la marca preservar_columnas (pestaña estándar) el esquema fijo
        # sigue mandando: el comportamiento anterior queda intacto.
        blob = _dossier_estilo_colpensiones()
        renombrado = renombrar_columnas_xlsx(blob, "TÍTULO", "RESUMEN")
        cfg = _cfg_custom()
        del cfg["preservar_columnas"]
        res = process_dossier(io.BytesIO(renombrado), {}, {}, ai_config=cfg)
        headers, ws, _ = _celdas_salida(res["output_data"])
        self.assertNotIn("ID", headers)
        self.assertNotIn("LINK", headers)
        self.assertIn("ID Noticia", headers)
        ws._parent.close()


if __name__ == "__main__":
    unittest.main()
