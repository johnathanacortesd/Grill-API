# ======================================
# Prueba de punta a punta del port: dossier XLSX real -> pipeline.process_dossier
# Sin API: el modelo va simulado (analyzer_tono_tema.llamar_llm).
# Verifica el contrato que el cliente ya usa: limpieza intacta, 4 columnas de
# analisis insertadas despues de 'revalorización', guarda del tono y uniformidad
# de etiquetas por grupo.
# ======================================
import io
import json
import re
import unittest

import openpyxl

import analyzer_tono_tema as A
import pipeline as P

MARCA = 'Universidad Simón Bolívar'
ALIAS = ['Unisimón']

CUBOS = ['Infraestructura educativa', 'Control y sanciones', 'Salud mental juvenil']

CABECERAS = [
    'ID Noticia', 'Fecha', 'Hora', 'Medio', 'Tipo de Medio', 'Sección - Programa', 'Región',
    'Título', 'Autor - Conductor', 'Nro. Pagina', 'Dimensión', 'Duración - Nro. Caracteres',
    'CPE', 'Tier', 'Audiencia', 'Valor de Nota', 'revalorización', 'resumen corto',
    'Link Nota', 'URL Nota', 'Resumen - Aclaracion', 'Link (Streaming - Imagen)',
    'Menciones - Empresa', 'Empresa rel.', 'ID duplicada',
]

NOTA_ENTREGA = (
    'La Universidad Simón Bolívar entregó el nuevo bloque de aulas de la sede norte. La obra '
    'beneficia a 1.200 estudiantes y fue financiada con recursos propios de la institución.'
)
NOTA_SANCION = (
    'La Contraloría sancionó a la Universidad Simón Bolívar por los sobrecostos en la entrega del '
    'bloque nuevo. El órgano de control cuestionó la demora de la institución en las obras.'
)
NOTA_SUICIDIO = (
    'Un informe nacional advierte que los intentos de suicidio se concentran en la población joven '
    'del país. El documento fue citado en la jornada de prevención de salud mental.'
)


def _filas():
    return [
        # 1 y 2: el mismo hecho publicado en dos medios -> un solo grupo
        {'ID Noticia': 1, 'Fecha': '2026-08-01', 'Hora': '08:00', 'Medio': 'El Heraldo',
         'Tipo de Medio': 'Internet', 'Región': 'Atlántico',
         'Título': 'La Universidad Simón Bolívar entrega nuevo bloque de aulas',
         'Resumen - Aclaracion': NOTA_ENTREGA, 'Link Nota': 'https://ejemplo.com/nota-1', 'URL Nota': 'https://ejemplo.com/nota-1',
         'Menciones - Empresa': MARCA, 'Empresa rel.': MARCA, 'CPE': '1.000.000', 'revalorización': '1.000.000'},
        {'ID Noticia': 2, 'Fecha': '2026-08-01', 'Hora': '09:30', 'Medio': 'Emisora Atlántico',
         'Tipo de Medio': 'Internet', 'Región': 'Atlántico',
         'Título': 'Universidad Simón Bolívar entregó su nuevo bloque de aulas',
         'Resumen - Aclaracion': NOTA_ENTREGA + ' La entrega se realizó en la sede norte.',
         'Link Nota': 'https://ejemplo.com/nota-2', 'URL Nota': 'https://ejemplo.com/nota-2',
         'Menciones - Empresa': MARCA, 'Empresa rel.': MARCA,
         'CPE': '2.000.000', 'revalorización': '2.000.000'},
        # 3: critica dirigida a la marca -> Negativo real
        {'ID Noticia': 3, 'Fecha': '2026-08-02', 'Hora': '10:00', 'Medio': 'El Tiempo',
         'Tipo de Medio': 'Internet', 'Región': 'Nacional',
         'Título': 'Sancionan a universidad por sobrecostos en obra',
         'Resumen - Aclaracion': NOTA_SANCION, 'Link Nota': 'https://ejemplo.com/nota-3', 'URL Nota': 'https://ejemplo.com/nota-3',
         'Menciones - Empresa': MARCA, 'Empresa rel.': MARCA, 'CPE': '3.000.000', 'revalorización': '3.000.000'},
        # 4: hecho tragico sin señalamiento -> el modelo dice Negativo y la guarda debe bajarlo
        {'ID Noticia': 4, 'Fecha': '2026-08-03', 'Hora': '11:00', 'Medio': 'Caracol Radio',
         'Tipo de Medio': 'Internet', 'Región': 'Nacional',
         'Título': 'Aumentan los intentos de suicidio en los jóvenes del país',
         'Resumen - Aclaracion': NOTA_SUICIDIO, 'Link Nota': 'https://ejemplo.com/nota-4', 'URL Nota': 'https://ejemplo.com/nota-4',
         'Menciones - Empresa': MARCA, 'Empresa rel.': MARCA, 'CPE': '4.000.000', 'revalorización': '4.000.000'},
        # 5: misma URL y misma mencion que la 1 -> duplicada
        {'ID Noticia': 5, 'Fecha': '2026-08-01', 'Hora': '12:00', 'Medio': 'El Heraldo',
         'Tipo de Medio': 'Internet', 'Región': 'Atlántico',
         'Título': 'La Universidad Simón Bolívar entrega nuevo bloque de aulas',
         'Resumen - Aclaracion': NOTA_ENTREGA, 'Link Nota': 'https://ejemplo.com/nota-1', 'URL Nota': 'https://ejemplo.com/nota-1',
         'Menciones - Empresa': MARCA, 'Empresa rel.': MARCA, 'CPE': '5.000.000', 'revalorización': '5.000.000'},
    ]


def _dossier_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(CABECERAS)
    for r in rows:
        ws.append([r.get(c) for c in CABECERAS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _modelo_simulado(cfg, mensajes, **kw):
    """Modelo simulado: reconoce cada prompt del motor y responde lo que corresponde."""
    prompt = ' '.join(m['content'] for m in mensajes)

    # 1) propuesta/consolidacion de la lista de Temas
    if 'CUBOS TEMATICOS' in prompt or 'LISTA FINAL' in prompt:
        return json.dumps({'cubos': list(CUBOS)}, ensure_ascii=False)

    # 2) asignacion de cubo a los grupos que no cazaron por reglas
    if 'Clasificas notas en cubos' in prompt:
        salida = []
        for sub in re.findall(r'SUB-TEMA: (.+)', prompt):
            s = sub.lower()
            cubo = ('Salud mental juvenil' if 'suicidio' in s else
                    'Control y sanciones' if 'sanci' in s or 'sobrecosto' in s else
                    'Infraestructura educativa')
            salida.append({'id': 0, 'cubo': cubo})
        ids = [int(x) for x in re.findall(r'GRUPO id=(\d+)', prompt)]
        for k, i in enumerate(ids):
            salida[k]['id'] = i
        return json.dumps({'resultados': salida}, ensure_ascii=False)

    # 3) reparacion (no deberia hacer falta: el modelo ya devuelve etiquetas validas)
    if 'Corrige SOLO estos sub-temas' in prompt:
        ids = [int(x) for x in re.findall(r'GRUPO id=(\d+)', prompt)]
        return json.dumps({'resultados': [{'id': i, 'sub_tema': 'Hecho informativo del dia',
                                           'tono': 'Neutro'} for i in ids]}, ensure_ascii=False)

    # 4) etiquetado del lote (sub-tema + tono por grupo)
    ids = [int(x) for x in re.findall(r'GRUPO id=(\d+)', prompt)]
    salida = []
    for i in ids:
        bloque = re.search(r'GRUPO id=%d\b.*?(?=\nGRUPO id=|\Z)' % i, prompt, re.S)
        texto = (bloque.group(0) if bloque else '').lower()
        if 'suicidio' in texto:
            # el modelo pequeño marca Negativo todo hecho tragico: la guarda debe corregirlo
            salida.append({'id': i, 'sub_tema': 'Aumento de intentos de suicidio', 'tono': 'Negativo'})
        elif 'sancion' in texto or 'sobrecosto' in texto:
            salida.append({'id': i, 'sub_tema': 'Sanción por sobrecostos en obra', 'tono': 'Negativo'})
        else:
            # programa/obra propia de la marca: el modelo la infravalora y dice Neutro;
            # la guarda positiva debe subirla a Positivo.
            salida.append({'id': i, 'sub_tema': 'Entrega de nuevo bloque de aulas', 'tono': 'Neutro'})
    return json.dumps({'resultados': salida}, ensure_ascii=False)


def _leer_salida(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    cab = list(filas[0])
    return cab, [dict(zip(cab, f)) for f in filas[1:]]


class TestPortPuntaAPunta(unittest.TestCase):

    def setUp(self):
        self._real = A.llamar_llm
        A.llamar_llm = _modelo_simulado
        self.res = P.process_dossier(
            io.BytesIO(_dossier_bytes(_filas())),
            region_map={}, internet_map={}, progress=None,
            ai_config={'enabled': True, 'brand': MARCA, 'aliases': ALIAS, 'voceros': [],
                       'criterio': list(A.CRITERIOS_TONO)[0],
                       'taxonomia': 'Automática según el archivo (recomendada)',
                       'cubos_objetivo': 8, 'votos': 2, 'permitir_cubos_nuevos': True,
                       'tam_lote': 10, 'workers': 1,
                       'api_key': 'fake', 'model': 'gpt-4.1-nano-2025-04-14'},
        )
        self.cab, self.filas = _leer_salida(self.res['output_data'])

    def tearDown(self):
        A.llamar_llm = self._real

    # --- contrato de salida -------------------------------------------------
    def test_columnas_de_analisis_quedan_despues_de_revalorizacion(self):
        i = self.cab.index('revalorización')
        self.assertEqual(self.cab[i + 1:i + 5],
                         ['Contexto analizado', 'Tono_IA', 'Tema_IA', 'Subtema_IA'])
        self.assertEqual(self.cab[i + 5], 'resumen corto')

    def test_limpieza_intacta_filas_y_duplicados(self):
        self.assertEqual(self.res['total_rows'], 5)
        self.assertEqual(self.res['unique_rows'], 4)
        self.assertEqual(len(self.filas), 5)
        # las columnas base siguen existiendo y en su orden
        for c in ('ID Noticia', 'Medio', 'Región', 'Link Nota', 'Menciones - Empresa'):
            self.assertIn(c, self.cab)

    def test_duplicada_conserva_su_marca(self):
        dup = [f for f in self.filas if f['ID Noticia'] == 5][0]
        self.assertEqual(dup['Tono_IA'], 'Duplicada')
        self.assertEqual(dup['Tema_IA'], '-')
        self.assertEqual(dup['Subtema_IA'], '-')
        self.assertEqual(dup['ID duplicada'], 1)

    # --- calidad del motor ---------------------------------------------------
    def test_todo_grupo_sale_con_etiquetas_validas(self):
        for f in self.filas:
            if f['Tono_IA'] == 'Duplicada':
                continue
            self.assertIn(f['Tono_IA'], A.TONOS)
            problemas = [p for p in A.validar(f['Subtema_IA'], f['Tono_IA'], [f['Título']])
                         if not p.startswith('revisar_anclaje')]
            self.assertEqual(problemas, [], '%s -> %s' % (f['Subtema_IA'], problemas))

    def test_ningun_tema_es_otros_y_hay_taxonomia_propia(self):
        for f in self.filas:
            self.assertNotEqual(A.nz(f['Tema_IA']), 'otros')
        self.assertTrue(self.res['analisis'].get('taxonomia'))
        for t in self.res['analisis']['taxonomia']:
            self.assertNotIn(A.nz(t), A.CUBO_PROHIBIDO)

    def test_notas_del_mismo_hecho_comparten_etiqueta(self):
        f1, f2 = [f for f in self.filas if f['ID Noticia'] in (1, 2)]
        self.assertEqual((f1['Tono_IA'], f1['Tema_IA'], f1['Subtema_IA']),
                         (f2['Tono_IA'], f2['Tema_IA'], f2['Subtema_IA']))
        self.assertEqual(f1['Tono_IA'], 'Positivo')

    def test_guarda_del_tono_baja_el_falso_negativo(self):
        # el modelo dijo Negativo en el hecho tragico; sin señalamiento dirigido debe quedar Neutro
        f4 = [f for f in self.filas if f['ID Noticia'] == 4][0]
        self.assertEqual(f4['Tono_IA'], 'Neutro')
        self.assertTrue(self.res['analisis'].get('tono_corregido_por_guarda'),
                        'la guarda debe reportar el grupo corregido')

    def test_critica_dirigida_a_la_marca_sigue_siendo_negativa(self):
        f3 = [f for f in self.filas if f['ID Noticia'] == 3][0]
        self.assertEqual(f3['Tono_IA'], 'Negativo')

    def test_guarda_positiva_sube_el_programa_propio(self):
        # el modelo dejo en Neutro la entrega de aulas de la propia marca; la guarda la sube
        f1 = [f for f in self.filas if f['ID Noticia'] == 1][0]
        self.assertEqual(f1['Tono_IA'], 'Positivo')
        self.assertTrue(self.res['analisis'].get('tono_subido_por_guarda'),
                        'la guarda positiva debe reportar el grupo corregido')

    def test_el_modelo_recibe_la_marca_y_el_modelo_configurado(self):
        capturado = {}

        prompts = []

        def _espia(cfg, mensajes, **kw):
            capturado['cfg'] = dict(cfg)
            prompts.append(' '.join(m['content'] for m in mensajes))
            return _modelo_simulado(cfg, mensajes, **kw)

        A.llamar_llm = _espia
        try:
            P.process_dossier(
                io.BytesIO(_dossier_bytes(_filas())), region_map={}, internet_map={}, progress=None,
                ai_config={'enabled': True, 'brand': MARCA, 'aliases': ALIAS, 'voceros': ['el rector'],
                           'criterio': list(A.CRITERIOS_TONO)[0],
                           'taxonomia': 'Automática según el archivo (recomendada)',
                           'cubos_objetivo': 8, 'votos': 1, 'permitir_cubos_nuevos': True,
                           'tam_lote': 10, 'workers': 1,
                           'api_key': 'llave-secreta-de-prueba', 'model': 'gpt-4.1-nano-2025-04-14'},
            )
        finally:
            A.llamar_llm = self._real
        self.assertEqual(capturado['cfg']['model'], 'gpt-4.1-nano-2025-04-14')
        self.assertEqual(capturado['cfg']['api_key'], 'llave-secreta-de-prueba')
        self.assertEqual(capturado['cfg']['base_url'], 'https://api.openai.com/v1')
        todo = '\n'.join(prompts)
        self.assertIn(MARCA, todo)
        self.assertIn('Unisimón', todo)
        self.assertIn('el rector', todo)
        self.assertIn('ALIAS Y FORMAS DE NOMBRARLA', todo)


if __name__ == '__main__':
    unittest.main()
