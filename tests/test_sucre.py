# ======================================
# Variante Sucre: Grill completo + actores (solo personas)
# ======================================
import ast
import io
import os
import sys
import unittest
from unittest.mock import MagicMock

import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from pipeline import BASE_OUTPUT_COLUMNS, KEY_MAP, PLAIN_HYPERLINK_COLUMNS
from sucre_analyzer import (
    COL_EXTERNOS,
    COL_INT_PROPIA,
    COL_MENCION_EXT,
    COL_PROPIOS,
    COL_TONO,
    LUCY_CANONICAL_NAME,
    SUCRE_OUTPUT_COLUMNS,
    analyze_article,
    enforce_people_only,
    format_extract,
    heuristic_analyze,
    is_entity_only_label,
    is_filler,
    lucy_intervenes,
    merge_analysis,
    recover_verbatim,
)
from sucre_pipeline import build_sample_xlsx, process_sucre_dossier


LUCY_BODY = (
    "La gobernadora Lucy García Montes entregó 200 becas universitarias en Sincelejo. "
    '"La educación es la prioridad de este gobierno", afirmó la mandataria. '
    "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social."
)

SENATOR_BODY = (
    "El senador Andrés Pérez señaló que la Gobernación de Sucre no ha ejecutado "
    "el presupuesto de vías rurales y pidió explicaciones a la mandataria."
)

SECRE_ENTITY_BODY = (
    "La Secretaría de Educación departamental emitió la Resolución 045 de 2026 "
    "para ampliar la cobertura escolar en los municipios del Golfo de Morrosquillo."
)

SECRETARY_PERSON_BODY = (
    "El secretario de Educación de Sucre, Carlos Méndez, anunció la ampliación "
    "de cupos escolares en Sincelejo y reiteró el compromiso de la Gobernación."
)

GOB_ONLY_BODY = (
    "La Gobernación de Sucre emitió un comunicado sobre el presupuesto de 2026 "
    "sin declaraciones de la mandataria ni de secretarios."
)

RENDON_BODY = (
    "Andrés Julián Rendón, Gobernador de Antioquia, afirmó que la Gobernación de Sucre "
    "puede replicar el modelo de vías terciarias del occidente antioqueño."
)

MADURO_BODY = (
    "A través de su cuenta de X, la gobernadora de Sucre, Lucy García Montes, "
    "aseguró que la captura de Nicolás Maduro debe convertirse en un punto de "
    "partida para reconstruir el rumbo en Venezuela y devolverle la esperanza a la gente."
)

MADURO_MID_CLAUSE = (
    "que la captura de Nicolás Maduro debe convertirse en un punto de partida"
)

EXT_MADURO_BODY = (
    "A través de su cuenta de X, el senador Andrés Pérez aseguró que la gobernadora "
    "de Sucre debe convertir la captura de Nicolás Maduro en un punto de partida "
    "para reconstruir el rumbo en Venezuela."
)


class VerbatimRulesTests(unittest.TestCase):
    def test_recover_exact_span(self):
        src = "La Gobernación de Sucre emitió un comunicado."
        self.assertEqual(recover_verbatim("La Gobernación de Sucre emitió un comunicado.", src), src)

    def test_rejects_paraphrase(self):
        src = "La gobernadora inauguró el hospital de Sampués."
        got = recover_verbatim("La mandataria abrió un nuevo centro hospitalario en Sampués.", src)
        self.assertEqual(got, "")

    def test_whitespace_flexible_recovers_source_slice(self):
        src = "Lucy  García\nanunció la obra."
        got = recover_verbatim("Lucy García anunció la obra.", src)
        self.assertIn("Lucy", got)
        self.assertIn("anunció la obra", got.replace("\n", " "))

    def test_filler_never_kept(self):
        self.assertTrue(is_filler("(sin actor externo)"))
        self.assertTrue(is_filler("N/A"))
        self.assertTrue(is_filler("ninguno"))
        self.assertTrue(is_filler("—"))
        self.assertFalse(is_filler("Lucy Inés García Montes, Gobernadora de Sucre"))


class ExtractClauseExpandTests(unittest.TestCase):
    def test_maduro_x_mid_clause_expands_to_sentence_start(self):
        got = format_extract(MADURO_MID_CLAUSE, MADURO_BODY)
        self.assertTrue(got.startswith("A través"))
        self.assertTrue(got[0].isupper())
        self.assertTrue(got.endswith("."))
        self.assertIn("Lucy García Montes", got)
        self.assertIn("aseguró que la captura de Nicolás Maduro", got)
        self.assertIn("devolverle la esperanza a la gente.", got)
        self.assertTrue(got[:-1] in MADURO_BODY or got in MADURO_BODY)

    def test_capitalizes_first_letter_and_appends_period(self):
        src = "la gobernadora de Sucre anunció la pavimentación de la vía"
        got = format_extract(src, src)
        self.assertEqual(got, "La gobernadora de Sucre anunció la pavimentación de la vía.")
        self.assertTrue(got[0].isupper())
        self.assertTrue(got.endswith("."))

    def test_keeps_existing_exclamation(self):
        src = "¡Lucy García Montes inauguró el hospital de Sampués!"
        got = format_extract("inauguró el hospital de Sampués", src)
        self.assertEqual(got, src)

    def test_heuristic_maduro_starts_at_full_clause(self):
        out = heuristic_analyze("Captura de Maduro", MADURO_BODY)
        extract = out[COL_INT_PROPIA]
        self.assertIn(LUCY_CANONICAL_NAME, out[COL_PROPIOS])
        self.assertTrue(extract.startswith("A través"))
        self.assertTrue(extract[0].isupper())
        self.assertTrue(extract.endswith("."))
        self.assertIn("cuenta de X", extract)

    def test_llm_mid_clause_is_expanded_on_merge(self):
        heuristic = heuristic_analyze("Captura de Maduro", MADURO_BODY)
        llm = {
            "tono": "Positivo",
            "nombre_cargo_propios": "Lucy Inés García Montes, Gobernadora de Sucre",
            "intervencion_propia": MADURO_MID_CLAUSE,
            "nombre_cargo_externos": "",
            "mencion_externa": "",
        }
        merged = merge_analysis(llm, heuristic, "Captura de Maduro", MADURO_BODY)
        extract = merged[COL_INT_PROPIA]
        self.assertTrue(extract.startswith("A través"))
        self.assertTrue(extract[0].isupper())
        self.assertTrue(extract.endswith("."))
        self.assertNotEqual(extract[:3].lower(), "que")
        self.assertIn("Lucy García Montes", extract)

    def test_external_mid_clause_also_expands(self):
        heuristic = heuristic_analyze("Senador sobre Maduro", EXT_MADURO_BODY)
        llm = {
            "tono": "Neutro",
            "nombre_cargo_propios": "",
            "intervencion_propia": "",
            "nombre_cargo_externos": "Andrés Pérez, senador",
            "mencion_externa": (
                "que la gobernadora de Sucre debe convertir la captura de Nicolás Maduro"
            ),
        }
        merged = merge_analysis(llm, heuristic, "Senador sobre Maduro", EXT_MADURO_BODY)
        extract = merged[COL_MENCION_EXT]
        self.assertIn("Andrés Pérez", merged[COL_EXTERNOS])
        self.assertTrue(extract.startswith("A través"))
        self.assertTrue(extract[0].isupper())
        self.assertTrue(extract.endswith("."))
        self.assertIn("senador Andrés Pérez", extract)


class PeopleOnlyLabelTests(unittest.TestCase):
    def test_forbidden_entities_are_not_people(self):
        for label in (
            "Secretaría de Educación departamental",
            "Gobernación de Sucre",
            "Ministerio del Interior",
            "Presidencia de la República",
        ):
            self.assertTrue(is_entity_only_label(label), label)

    def test_person_plus_role_is_not_entity(self):
        self.assertFalse(is_entity_only_label("Andrés Julián Rendón, Gobernador de Antioquia"))
        self.assertFalse(is_entity_only_label("Lucy Inés García Montes, Gobernadora de Sucre"))


class HeuristicExtractTests(unittest.TestCase):
    def test_lucy_own_intervention_is_literal_from_body(self):
        out = heuristic_analyze("Gobernadora entregó becas", LUCY_BODY)
        self.assertEqual(out[COL_TONO], "Positivo")
        self.assertIn(LUCY_CANONICAL_NAME, out[COL_PROPIOS])
        self.assertIn("Gobernadora", out[COL_PROPIOS])
        extract = out[COL_INT_PROPIA]
        self.assertTrue(extract)
        self.assertIn(extract[:40], LUCY_BODY)
        self.assertIn("entregó 200 becas", extract)
        self.assertNotIn("sin actor", extract.lower())
        self.assertIn("Ricardo Hernández", out[COL_EXTERNOS])
        self.assertIn("Alcalde", out[COL_EXTERNOS])
        ext = out[COL_MENCION_EXT]
        self.assertTrue(ext)
        self.assertIn("alcalde", ext.lower())
        self.assertIn(ext[:30], LUCY_BODY)

    def test_external_senator_no_own_agency(self):
        out = heuristic_analyze("Senador cuestiona vías", SENATOR_BODY)
        self.assertEqual(out[COL_TONO], "Negativo")
        self.assertEqual(out[COL_INT_PROPIA], "")
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertIn("Andrés Pérez", out[COL_EXTERNOS])
        self.assertIn("senador", out[COL_EXTERNOS].lower())
        self.assertIn(out[COL_MENCION_EXT][:40], SENATOR_BODY)
        self.assertNotEqual(out[COL_EXTERNOS].lower(), "(sin actor externo)")

    def test_secretaria_entity_is_not_propio(self):
        out = heuristic_analyze("Resolución de cobertura", SECRE_ENTITY_BODY)
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertEqual(out[COL_INT_PROPIA], "")
        self.assertNotIn("Secretaría de Educación", out[COL_PROPIOS])
        self.assertEqual(out[COL_EXTERNOS], "")
        self.assertEqual(out[COL_MENCION_EXT], "")

    def test_named_secretary_is_propio_person(self):
        out = heuristic_analyze("Ampliación de cupos", SECRETARY_PERSON_BODY)
        self.assertIn("Carlos Méndez", out[COL_PROPIOS])
        self.assertIn("Secretario", out[COL_PROPIOS])
        self.assertIn("anunció la ampliación", out[COL_INT_PROPIA])
        self.assertIn(out[COL_INT_PROPIA][:40], SECRETARY_PERSON_BODY)
        self.assertNotIn("Gobernación de Sucre", out[COL_PROPIOS])

    def test_gobernacion_without_lucy_intervention_leaves_propios_empty(self):
        out = heuristic_analyze("Comunicado presupuestal", GOB_ONLY_BODY)
        self.assertFalse(lucy_intervenes("Comunicado presupuestal", GOB_ONLY_BODY))
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertEqual(out[COL_INT_PROPIA], "")
        self.assertNotIn("Lucy", out[COL_PROPIOS])

    def test_external_governor_requires_name_and_role(self):
        out = heuristic_analyze("Gestión vial", RENDON_BODY)
        self.assertIn("Andrés Julián Rendón", out[COL_EXTERNOS])
        self.assertIn("Gobernador de Antioquia", out[COL_EXTERNOS])
        self.assertIn(out[COL_MENCION_EXT][:40], RENDON_BODY)
        self.assertEqual(out[COL_PROPIOS], "")

    def test_external_without_sucre_or_lucy_mention_is_empty(self):
        body = (
            "La Gobernación de Sucre presentó el plan departamental de vías. "
            "En Bogotá, el senador Andrés Pérez opinó que el Gobierno nacional "
            "debe acelerar la reforma a la salud y no habló de Sincelejo."
        )
        out = heuristic_analyze("Plan vial y debate nacional", body)
        self.assertEqual(out[COL_EXTERNOS], "")
        self.assertEqual(out[COL_MENCION_EXT], "")

    def test_external_with_sucre_mention_fills_person_and_extract(self):
        out = heuristic_analyze("Senador cuestiona vías", SENATOR_BODY)
        self.assertIn("Andrés Pérez", out[COL_EXTERNOS])
        self.assertIn("senador", out[COL_EXTERNOS].lower())
        extract = out[COL_MENCION_EXT]
        self.assertTrue(extract)
        self.assertIn(extract[:40], SENATOR_BODY)
        self.assertIn("Gobernación de Sucre", extract)
        self.assertIn("Andrés Pérez", extract)

    def test_entity_only_external_is_empty(self):
        body = (
            "El Ministerio del Interior cuestionó a la Gobernación de Sucre "
            "por el rezago en la ejecución de vías."
        )
        out = heuristic_analyze("Cuestionamiento nacional", body)
        self.assertEqual(out[COL_EXTERNOS], "")
        self.assertEqual(out[COL_MENCION_EXT], "")
        self.assertNotIn("Ministerio", out[COL_EXTERNOS])

    def test_no_mention_yields_empty_strings(self):
        body = "El Ideam publicó el pronóstico de lluvias para la región Caribe."
        out = heuristic_analyze("Boletín climático", body)
        self.assertEqual(out[COL_TONO], "Neutro")
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertEqual(out[COL_INT_PROPIA], "")
        self.assertEqual(out[COL_EXTERNOS], "")
        self.assertEqual(out[COL_MENCION_EXT], "")

    def test_alias_gobernadora_de_sucre(self):
        body = "La gobernadora de Sucre anunció la pavimentación de la vía Sincelejo-Sampués."
        out = heuristic_analyze("Pavimentación vial", body)
        self.assertIn(LUCY_CANONICAL_NAME, out[COL_PROPIOS])
        self.assertIn("anunció la pavimentación", out[COL_INT_PROPIA])
        self.assertIn(out[COL_INT_PROPIA], body)

    def test_lucy_montes_variant(self):
        body = "Lucy Montes inauguró el hospital de Sampués ante la comunidad."
        out = heuristic_analyze("Hospital de Sampués", body)
        self.assertIn(LUCY_CANONICAL_NAME, out[COL_PROPIOS])
        self.assertIn("inauguró el hospital", out[COL_INT_PROPIA])


class LlmMergeTests(unittest.TestCase):
    def test_paraphrased_llm_extract_falls_back_to_heuristic(self):
        heuristic = heuristic_analyze("Becas", LUCY_BODY)
        llm = {
            COL_TONO: "Positivo",
            COL_PROPIOS: "Lucy Inés García Montes, Gobernadora de Sucre",
            COL_INT_PROPIA: "La mandataria entregó un paquete de apoyos académicos.",
            COL_EXTERNOS: "",
            COL_MENCION_EXT: "(sin actor externo)",
        }
        merged = merge_analysis(llm, heuristic, "Becas", LUCY_BODY)
        self.assertIn("entregó 200 becas", merged[COL_INT_PROPIA])
        self.assertNotIn("paquete de apoyos", merged[COL_INT_PROPIA])
        self.assertNotIn("sin actor", merged[COL_MENCION_EXT].lower())

    def test_verbatim_llm_span_is_kept(self):
        span = "La gobernadora Lucy García Montes entregó 200 becas universitarias en Sincelejo."
        heuristic = heuristic_analyze("Becas", LUCY_BODY)
        llm = {
            "tono": "Positivo",
            "nombre_cargo_propios": "Lucy Inés García Montes, Gobernadora de Sucre",
            "intervencion_propia": span,
            "nombre_cargo_externos": "Ricardo Hernández, alcalde",
            "mencion_externa": "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social.",
        }
        merged = merge_analysis(llm, heuristic, "Becas", LUCY_BODY)
        self.assertEqual(merged[COL_INT_PROPIA], span)
        self.assertEqual(
            merged[COL_MENCION_EXT],
            "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social.",
        )

    def test_llm_entity_propio_is_stripped(self):
        heuristic = heuristic_analyze("Resolución", SECRE_ENTITY_BODY)
        llm = {
            COL_PROPIOS: "Secretaría de Educación departamental",
            COL_INT_PROPIA: "La Secretaría de Educación departamental emitió la Resolución 045 de 2026",
            COL_EXTERNOS: "Ministerio del Interior",
            COL_MENCION_EXT: "",
        }
        merged = merge_analysis(llm, heuristic, "Resolución", SECRE_ENTITY_BODY)
        self.assertEqual(merged[COL_PROPIOS], "")
        self.assertEqual(merged[COL_INT_PROPIA], "")
        self.assertEqual(merged[COL_EXTERNOS], "")

    def test_llm_cannot_invent_lucy_without_intervention(self):
        heuristic = heuristic_analyze("Comunicado", GOB_ONLY_BODY)
        llm = {
            COL_PROPIOS: "Lucy Inés García Montes, Gobernadora de Sucre",
            COL_INT_PROPIA: "La Gobernación de Sucre emitió un comunicado sobre el presupuesto de 2026",
            COL_EXTERNOS: "",
            COL_MENCION_EXT: "",
        }
        merged = merge_analysis(llm, heuristic, "Comunicado", GOB_ONLY_BODY)
        self.assertEqual(merged[COL_PROPIOS], "")
        self.assertNotIn("Lucy", merged[COL_PROPIOS])

    def test_analyze_article_without_client_is_heuristic(self):
        out = analyze_article("Becas", LUCY_BODY, client=None)
        self.assertEqual(out[COL_TONO], "Positivo")
        self.assertTrue(out[COL_INT_PROPIA])

    def test_llm_external_without_sucre_lucy_in_extract_is_dropped(self):
        body = (
            "La gobernadora Lucy García anunció becas en Sincelejo. "
            "El senador Andrés Pérez dijo que el Congreso debe votar la reforma pensional."
        )
        heuristic = heuristic_analyze("Becas y reforma", body)
        llm = {
            "tono": "Positivo",
            "nombre_cargo_propios": "Lucy Inés García Montes, Gobernadora de Sucre",
            "intervencion_propia": "La gobernadora Lucy García anunció becas en Sincelejo.",
            "nombre_cargo_externos": "Andrés Pérez, senador",
            "mencion_externa": (
                "El senador Andrés Pérez dijo que el Congreso debe votar la reforma pensional."
            ),
        }
        merged = merge_analysis(llm, heuristic, "Becas y reforma", body)
        self.assertEqual(merged[COL_EXTERNOS], "")
        self.assertEqual(merged[COL_MENCION_EXT], "")

    def test_llm_bad_extract_recovers_span_that_mentions_gobernacion(self):
        body = (
            "El senador Andrés Pérez saludó a los periodistas en el recinto. "
            "Más tarde el senador Andrés Pérez señaló que la Gobernación de Sucre "
            "no ha ejecutado el presupuesto de vías rurales."
        )
        heuristic = heuristic_analyze("Senador", body)
        llm = {
            "tono": "Negativo",
            "nombre_cargo_propios": "",
            "intervencion_propia": "",
            "nombre_cargo_externos": "Andrés Pérez, senador",
            "mencion_externa": "El senador Andrés Pérez saludó a los periodistas en el recinto.",
        }
        merged = merge_analysis(llm, heuristic, "Senador", body)
        self.assertIn("Andrés Pérez", merged[COL_EXTERNOS])
        self.assertIn("Gobernación de Sucre", merged[COL_MENCION_EXT])
        self.assertNotIn("periodistas", merged[COL_MENCION_EXT])

    def test_enforce_drops_presidencia_alone(self):
        raw = {
            COL_TONO: "Neutro",
            COL_PROPIOS: "Presidencia de la República",
            COL_INT_PROPIA: "",
            COL_EXTERNOS: "Presidencia de la República",
            COL_MENCION_EXT: "",
        }
        out = enforce_people_only(raw, "Nota", "La Presidencia de la República mencionó a Sucre.")
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertEqual(out[COL_EXTERNOS], "")


class PipelineXlsxTests(unittest.TestCase):
    def test_sample_processing_keeps_grill_columns_and_adds_sucre(self):
        result = process_sucre_dossier(
            io.BytesIO(build_sample_xlsx()),
            region_map={},
            internet_map={},
            ai_config={"enabled": False},
        )
        self.assertGreaterEqual(result["total_rows"], 7)
        self.assertIn("unique_rows", result)
        self.assertIn("duplicates", result)
        self.assertTrue(result["output_filename"].startswith("Dossier_Gobernacion_de_Sucre"))

        bio = io.BytesIO(result["output_data"])
        wb = load_workbook(bio, read_only=True, data_only=True)
        ws = wb["Resultado"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for col in BASE_OUTPUT_COLUMNS:
            self.assertIn(col, headers)
        for col in SUCRE_OUTPUT_COLUMNS:
            self.assertIn(col, headers)
        self.assertNotIn("Tono", headers)
        self.assertNotIn("Tono_IA", headers)

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        by_header = [{headers[i]: row[i] for i in range(len(headers))} for row in rows]

        lucy_row = next(r for r in by_header if "becas" in str(r.get("Título") or "").lower())
        body = str(lucy_row["Resumen - Aclaracion"] or lucy_row.get("resumen corto") or "")
        own = str(lucy_row[COL_INT_PROPIA] or "")
        self.assertTrue(own)
        self.assertIn(own[:40], body)
        self.assertIn("Lucy", str(lucy_row[COL_PROPIOS] or ""))
        ext = str(lucy_row[COL_MENCION_EXT] or "")
        self.assertTrue(ext)
        self.assertIn(ext[:30], body)

        sen_row = next(r for r in by_header if "senador" in str(r.get("Título") or "").lower())
        self.assertIn(str(sen_row[COL_MENCION_EXT] or "")[:40], str(sen_row["Resumen - Aclaracion"] or ""))
        self.assertFalse(sen_row[COL_INT_PROPIA])
        self.assertFalse(sen_row[COL_PROPIOS])
        self.assertIn("Andrés Pérez", str(sen_row[COL_EXTERNOS] or ""))

        secre = next(r for r in by_header if "resolución" in str(r.get("Título") or "").lower() or "resolucion" in str(r.get("Título") or "").lower())
        self.assertEqual(secre[COL_PROPIOS] or "", "")
        self.assertEqual(secre[COL_INT_PROPIA] or "", "")

        clima = next(r for r in by_header if "climático" in str(r.get("Título") or "").lower() or "climatico" in str(r.get("Título") or "").lower())
        self.assertEqual(clima[COL_PROPIOS] or "", "")
        self.assertEqual(clima[COL_INT_PROPIA] or "", "")
        self.assertEqual(clima[COL_EXTERNOS] or "", "")
        self.assertEqual(clima[COL_MENCION_EXT] or "", "")

        sec_person = next(r for r in by_header if "cupos" in str(r.get("Título") or "").lower())
        self.assertIn("Carlos Méndez", str(sec_person[COL_PROPIOS] or ""))
        self.assertIn("anunció", str(sec_person[COL_INT_PROPIA] or ""))

        gob_only = next(r for r in by_header if "presupuestal" in str(r.get("Título") or "").lower())
        self.assertEqual(gob_only[COL_PROPIOS] or "", "")
        self.assertNotIn("Lucy", str(gob_only[COL_PROPIOS] or ""))

        rendon = next(r for r in by_header if "antioquia" in str(r.get("Título") or "").lower())
        self.assertIn("Andrés Julián Rendón", str(rendon[COL_EXTERNOS] or ""))
        self.assertIn("Gobernador de Antioquia", str(rendon[COL_EXTERNOS] or ""))
        wb.close()

    def test_link_style_columns_still_plain_hyperlinks(self):
        result = process_sucre_dossier(
            io.BytesIO(build_sample_xlsx()),
            region_map={},
            internet_map={},
            ai_config={"enabled": False},
        )
        wb = load_workbook(io.BytesIO(result["output_data"]))
        ws = wb["Resultado"]
        headers = [c.value for c in ws[1]]
        for col_name in ("Link Nota", "Link (Streaming - Imagen)"):
            self.assertIn(col_name, PLAIN_HYPERLINK_COLUMNS)
            self.assertIn(col_name, headers)
        wb.close()

    def test_pkl_path_keeps_grill_ai_columns_and_sucre_actors(self):
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.naive_bayes import MultinomialNB
        from sklearn.pipeline import Pipeline
        import joblib

        pipe = Pipeline([("tfidf", TfidfVectorizer()), ("clf", MultinomialNB())])
        pipe.fit(
            ["felicitaciones alianza beneficio", "denuncia grave corrupcion", "informe reunion cifras"],
            [1, -1, 0],
        )
        buf = io.BytesIO()
        joblib.dump(pipe, buf)
        result = process_sucre_dossier(
            io.BytesIO(build_sample_xlsx()),
            region_map={},
            internet_map={},
            ai_config={"enabled": False, "tone_pkl_bytes": buf.getvalue()},
        )
        df = pd.read_excel(io.BytesIO(result["output_data"]))
        self.assertIn("Tono_IA", df.columns)
        self.assertIn("Tema_IA", df.columns)
        self.assertIn("Subtema_IA", df.columns)
        for col in SUCRE_OUTPUT_COLUMNS:
            self.assertIn(col, df.columns)
        lucy = df[df["Título"].astype(str).str.contains("becas", case=False, na=False)].iloc[0]
        self.assertIn("Lucy", str(lucy[COL_PROPIOS]))

    def test_odd_columns_do_not_crash(self):
        df = pd.DataFrame({"Foo": [1], "Bar": ["x"]})
        buf = io.BytesIO()
        df.to_excel(buf, index=False, engine="openpyxl")
        result = process_sucre_dossier(
            io.BytesIO(buf.getvalue()),
            region_map={},
            internet_map={},
            ai_config={"enabled": False},
        )
        df_out = pd.read_excel(io.BytesIO(result["output_data"]))
        for col in BASE_OUTPUT_COLUMNS:
            self.assertIn(col, df_out.columns)
        for col in SUCRE_OUTPUT_COLUMNS:
            self.assertIn(col, df_out.columns)


class GrillIsolationTests(unittest.TestCase):
    def test_grill_output_shape_unchanged(self):
        self.assertNotIn(COL_INT_PROPIA, BASE_OUTPUT_COLUMNS)
        self.assertNotIn(COL_MENCION_EXT, BASE_OUTPUT_COLUMNS)
        self.assertNotIn(COL_PROPIOS, BASE_OUTPUT_COLUMNS)
        self.assertIn("Título", BASE_OUTPUT_COLUMNS)
        self.assertEqual(KEY_MAP["titulo"], "Título")

    def test_grill_app_does_not_import_sucre(self):
        with open(os.path.join(ROOT, "app.py"), encoding="utf-8") as fh:
            tree = ast.parse(fh.read())
        imported = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.extend(a.name for a in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.append(node.module)
        self.assertNotIn("sucre_analyzer", imported)
        self.assertNotIn("sucre_pipeline", imported)
        self.assertNotIn("app_sucre", imported)

    def test_sucre_app_is_full_grill_clone(self):
        with open(os.path.join(ROOT, "app_sucre.py"), encoding="utf-8") as fh:
            src = fh.read()
        self.assertIn("pkl_tono", src)
        self.assertIn("pkl_tema", src)
        self.assertIn("REGIONES_CSV_URL", src)
        self.assertIn("INTERNET_CSV_URL", src)
        self.assertIn("process_sucre_dossier", src)
        self.assertIn("Activar análisis reputacional", src)
        self.assertIn("Gobernación de Sucre", src)
        self.assertIn("app_sucre.py", src)
        self.assertNotIn("from pipeline import process_dossier", src)


class MockedLlmPathTests(unittest.TestCase):
    def test_openai_path_uses_literal_body_span(self):
        span = "La gobernadora Lucy García Montes entregó 200 becas universitarias en Sincelejo."
        payload = {
            "tono": "Positivo",
            "nombre_cargo_propios": "Lucy Inés García Montes, Gobernadora de Sucre",
            "intervencion_propia": span,
            "nombre_cargo_externos": "Ricardo Hernández, alcalde",
            "mencion_externa": "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social.",
        }
        fake_resp = MagicMock()
        fake_resp.choices = [MagicMock()]
        fake_resp.choices[0].message.content = __import__("json").dumps(payload)
        client = MagicMock()
        client.chat.completions.create.return_value = fake_resp
        out = analyze_article("Becas", LUCY_BODY, client=client, model="gpt-test")
        self.assertEqual(out[COL_INT_PROPIA], span)
        self.assertEqual(out[COL_TONO], "Positivo")
        client.chat.completions.create.assert_called_once()


if __name__ == "__main__":
    unittest.main()
