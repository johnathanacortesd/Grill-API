# -*- coding: utf-8 -*-
"""v4.32: modo solo-PKL (IA desactivada) no pierde la columna Tema_IA ni
exporta Subtema_IA vacío.

Bug reportado 2026-09-28 (pestaña «Columnas personalizadas»): con PKL de
tono + PKL de tema subidos, IA desactivada y el radio de Tema_IA en su
valor por defecto («Solo Tono_IA + Subtema_IA»), el xlsx salía con Tono_IA
(lleno), Subtema_IA (vacío, "-") y SIN columna Tema_IA, aunque el PKL de
tema había clasificado todo.

Causa: en `process_dossier`, el forzado de incluir_tema por PKL de tema
(`_ai_extra_con_pkl`) solo se aplicaba en el camino con IA; en el camino
solo-PKL se leía el ai_config crudo y la columna Tema_IA se eliminaba del
export. Además, sin IA el subtema nunca se genera (ningún PKL lo produce),
así que la columna salía vacía.

Fix (pipeline.py, general para ambas pestañas):
- `ai_extra = _ai_extra_con_pkl(ai_config, theme_model)` se calcula antes
  de la bifurcación y se usa para `cols_to_export` en ambos caminos: con
  PKL de tema siempre hay columna Tema_IA.
- `output_columns_for_export(..., include_subtema=False)` omite Subtema_IA;
  `process_dossier` pasa `include_subtema=has_ai` (sin IA no hay subtema).
"""
from __future__ import annotations

import io
import os
import sys
import types
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Stub del paquete `openai` (mismo patrón que test_sin_tema.py): pipeline lo
# importa vía ai_analyzer y ningún test de este archivo hace llamadas reales.
_openai_stub = types.ModuleType('openai')
_openai_stub.OpenAI = object
sys.modules.setdefault('openai', _openai_stub)

from openpyxl import Workbook, load_workbook

from pipeline import (
    _ai_extra_con_pkl,
    output_columns_for_export,
    process_dossier,
)


def _joblib_bytes(estimator):
    import io
    import joblib
    buf = io.BytesIO()
    joblib.dump(estimator, buf)
    return buf.getvalue()


def _fit_tono():
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.naive_bayes import MultinomialNB
    from sklearn.pipeline import make_pipeline
    texts = ["colpensiones excelente servicio bueno",
             "colpensiones pésimo servicio malo terrible",
             "colpensiones informa trámites horarios"]
    labels = ["Positivo", "Negativo", "Neutro"]
    clf = make_pipeline(TfidfVectorizer(), MultinomialNB())
    clf.fit(texts, labels)
    return clf


def _fit_tema():
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.naive_bayes import MultinomialNB
    from sklearn.pipeline import make_pipeline
    texts = ["pago de mesadas pensiones dinero",
             "trámite afiliación formulario ventanilla",
             "reforma ley congreso normativa"]
    labels = ["Entorno Pensional", "Trámites y Servicios", "Normativa y Regulación"]
    clf = make_pipeline(TfidfVectorizer(), MultinomialNB())
    clf.fit(texts, labels)
    return clf


def _dossier_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(["Título", "Resumen - Aclaracion", "Medio", "Fecha"])
    ws.append(["Colpensiones paga las mesadas puntualmente",
               "pago de mesadas pensiones dinero", "Prensa", "2026-09-01"])
    ws.append(["Trámite de afiliación en ventanilla",
               "trámite afiliación formulario ventanilla", "Prensa", "2026-09-02"])
    ws.append(["Nueva reforma de la ley pensional",
               "reforma ley congreso normativa", "Prensa", "2026-09-03"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ai_config_solo_pkl(tone_bytes, theme_bytes):
    return {
        "enabled": False,  # IA desactivada: camino solo-PKL
        "brand": "Colpensiones",
        "aliases": [],
        "voceros": [],
        "criterio": "General",
        "criterio_texto": "",
        "incluir_tema": False,  # radio en su defecto: "Solo Tono_IA + Subtema_IA"
        "incluir_prominencia": False,
        "tone_pkl_bytes": tone_bytes,
        "theme_pkl_bytes": theme_bytes,
    }


def _columnas_salida(xlsx_bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.active
    cols = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return cols


class TestSoloPklColumnasV432(unittest.TestCase):
    def test_extra_fuerza_tema_con_pkl(self):
        extra = _ai_extra_con_pkl({"incluir_tema": False}, object())
        self.assertTrue(extra["incluir_tema"])
        extra2 = _ai_extra_con_pkl({"incluir_tema": False}, None)
        self.assertFalse(extra2["incluir_tema"])

    def test_output_columns_sin_subtema(self):
        cols = output_columns_for_export(
            include_ai=True, include_tema=True, include_subtema=False)
        self.assertIn("Tono_IA", cols)
        self.assertIn("Tema_IA", cols)
        self.assertNotIn("Subtema_IA", cols)

    def test_output_columns_default_intacto(self):
        cols = output_columns_for_export(include_ai=True)
        self.assertIn("Tono_IA", cols)
        self.assertIn("Tema_IA", cols)
        self.assertIn("Subtema_IA", cols)

    def test_solo_pkl_tono_y_tema_columnas(self):
        tone_bytes = _joblib_bytes(_fit_tono())
        theme_bytes = _joblib_bytes(_fit_tema())
        res = process_dossier(
            io.BytesIO(_dossier_bytes()), {}, {},
            ai_config=_ai_config_solo_pkl(tone_bytes, theme_bytes))
        cols = _columnas_salida(res["output_data"])
        # El PKL de tema manda: hay columna Tema_IA aunque el radio dijera no.
        self.assertIn("Tono_IA", cols)
        self.assertIn("Tema_IA", cols)
        # Sin IA no hay subtema: la columna no se exporta (antes salía vacía).
        self.assertNotIn("Subtema_IA", cols)

    def test_solo_pkl_tono_sin_tema(self):
        tone_bytes = _joblib_bytes(_fit_tono())
        res = process_dossier(
            io.BytesIO(_dossier_bytes()), {}, {},
            ai_config=_ai_config_solo_pkl(tone_bytes, None))
        cols = _columnas_salida(res["output_data"])
        self.assertIn("Tono_IA", cols)
        self.assertNotIn("Tema_IA", cols)
        self.assertNotIn("Subtema_IA", cols)


if __name__ == "__main__":
    unittest.main()
