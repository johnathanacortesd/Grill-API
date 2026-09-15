# ======================================
# Columnas del Excel de salida y formato de Resumen - Aclaracion
# ======================================
import io
import os
import sys
import unittest

import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from pipeline import (  # noqa: E402
    AI_OUTPUT_COLUMNS,
    BASE_OUTPUT_COLUMNS,
    KEY_MAP,
    corregir_texto,
    generate_output_excel,
    output_columns_for_export,
    process_dossier,
    terminar_en_punto,
)

REMOVED_EXPORT_COLUMNS = ("revalorización", "resumen corto")


class OutputColumnOrderTests(unittest.TestCase):
    def test_base_columns_omit_resumen_corto_and_revalorizacion(self):
        for col in REMOVED_EXPORT_COLUMNS:
            self.assertNotIn(col, BASE_OUTPUT_COLUMNS)
        self.assertNotIn("revalorizacion", [c.lower() for c in BASE_OUTPUT_COLUMNS])

    def test_ai_export_puts_contexto_analizado_last(self):
        cols = output_columns_for_export(include_ai=True)
        self.assertEqual(cols[-4:], list(AI_OUTPUT_COLUMNS))
        self.assertEqual(cols[-1], "Contexto analizado")
        for col in REMOVED_EXPORT_COLUMNS:
            self.assertNotIn(col, cols)

    def test_non_ai_export_matches_base_and_omits_ai_cols(self):
        cols = output_columns_for_export(include_ai=False)
        self.assertEqual(cols, list(BASE_OUTPUT_COLUMNS))
        self.assertNotIn("Contexto analizado", cols)
        self.assertNotIn("Tono_IA", cols)


class ResumenAclaracionFormatTests(unittest.TestCase):
    def test_terminar_en_punto_replaces_ellipsis(self):
        self.assertEqual(terminar_en_punto("Texto..."), "Texto.")
        self.assertEqual(terminar_en_punto("Texto."), "Texto.")
        self.assertEqual(terminar_en_punto("Texto"), "Texto.")

    def test_corregir_texto_ends_with_single_period_not_ellipsis(self):
        out = corregir_texto("La universidad entrega aulas")
        self.assertTrue(out.endswith("."))
        self.assertFalse(out.endswith("..."))
        self.assertEqual(out, "La universidad entrega aulas.")

        already = corregir_texto("La nota ya termina.")
        self.assertEqual(already, "La nota ya termina.")

        ellipsis = corregir_texto("La nota termina con puntos...")
        self.assertEqual(ellipsis, "La nota termina con puntos.")
        self.assertFalse(ellipsis.endswith("..."))

    def test_exported_resumen_aclaracion_ends_with_period(self):
        rows = [
            {
                "ID Noticia": 1,
                "Título": "Nota de prueba",
                "Resumen - Aclaracion": "La universidad entrega aulas",
                "resumen corto": "no debe exportarse",
                "revalorización": 1000,
            }
        ]
        data = generate_output_excel(rows, KEY_MAP)
        wb = load_workbook(io.BytesIO(data))
        ws = wb["Resultado"]
        headers = [cell.value for cell in ws[1]]
        for col in REMOVED_EXPORT_COLUMNS:
            self.assertNotIn(col, headers)
        idx = headers.index("Resumen - Aclaracion") + 1
        value = ws.cell(row=2, column=idx).value
        self.assertTrue(str(value).endswith("."))
        self.assertFalse(str(value).endswith("..."))
        wb.close()

    def test_resumen_corto_input_is_read_but_not_exported(self):
        df = pd.DataFrame(
            {
                "NoticiaId": [1],
                "Fecha": ["01/01/2026"],
                "Tipo de Medio": ["internet"],
                "Título": ["Nota de prueba"],
                "resumen corto": ["La universidad entrega aulas en la sede norte"],
                "Empresa rel.": ["Marca Demo"],
            }
        )
        buf = io.BytesIO()
        df.to_excel(buf, index=False, engine="openpyxl")
        result = process_dossier(buf.getvalue(), region_map={}, internet_map={}, ai_config=None)
        out = pd.read_excel(io.BytesIO(result["output_data"]))
        self.assertNotIn("resumen corto", out.columns)
        self.assertNotIn("revalorización", out.columns)
        resumen = str(out.loc[0, "Resumen - Aclaracion"])
        self.assertIn("universidad entrega aulas", resumen)
        self.assertTrue(resumen.endswith("."))
        self.assertFalse(resumen.endswith("..."))


if __name__ == "__main__":
    unittest.main()
