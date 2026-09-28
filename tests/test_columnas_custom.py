"""v4.27: pestaña "Columnas personalizadas" — lectura, sugerencia y renombre
de las columnas de título y cuerpo (CuerpoEs) sin tocar el pipeline."""
import io
import os
import sys
import types
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Stub del paquete `openai`: no instalable en este entorno (mismo patrón que
# test_sin_tema.py). Ningún test de este archivo hace llamadas reales a la API.
_openai_stub = types.ModuleType('openai')
_openai_stub.OpenAI = object
sys.modules.setdefault('openai', _openai_stub)

from pipeline import (leer_columnas_xlsx, renombrar_columnas_xlsx, sugerir_columna,
                      construir_ai_config_custom)


def _xlsx_bytes(headers, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def _headers_of(blob):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        return [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    finally:
        wb.close()


def _row2_of(blob):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        return [c.value for c in next(wb.active.iter_rows(min_row=2, max_row=2))]
    finally:
        wb.close()


class TestColumnasCustom(unittest.TestCase):
    def test_lee_columnas(self):
        blob = _xlsx_bytes(["Headline", "Body", "Medio"], [["t1", "b1", "m1"]])
        cols = leer_columnas_xlsx(blob)
        self.assertEqual([lbl for lbl, _ in cols], ["Headline", "Body", "Medio"])
        self.assertEqual([idx for _, idx in cols], [1, 2, 3])

    def test_duplicados_reciben_etiqueta_unica(self):
        blob = _xlsx_bytes(["A", "A", "B"], [["1", "2", "3"]])
        cols = leer_columnas_xlsx(blob)
        self.assertEqual([lbl for lbl, _ in cols], ["A", "A (2)", "B"])

    def test_renombra_y_preserva_datos(self):
        blob = _xlsx_bytes(["Headline", "Body", "Medio"], [["t1", "b1", "m1"]])
        out = renombrar_columnas_xlsx(blob, "Headline", "Body")
        self.assertEqual(_headers_of(out), ["Título", "Resumen - Aclaracion", "Medio"])
        self.assertEqual(_row2_of(out), ["t1", "b1", "m1"])

    def test_aparta_columna_que_ya_tenia_el_nombre(self):
        blob = _xlsx_bytes(["Título", "Headline", "Body"], [["x", "t1", "b1"]])
        out = renombrar_columnas_xlsx(blob, "Headline", "Body")
        self.assertEqual(_headers_of(out),
                         ["Título (original)", "Título", "Resumen - Aclaracion"])
        self.assertEqual(_row2_of(out), ["x", "t1", "b1"])

    def test_misma_columna_da_error(self):
        blob = _xlsx_bytes(["A", "B"], [["1", "2"]])
        with self.assertRaises(ValueError):
            renombrar_columnas_xlsx(blob, "A", "A")

    def test_columna_inexistente_da_error(self):
        blob = _xlsx_bytes(["A", "B"], [["1", "2"]])
        with self.assertRaises(ValueError):
            renombrar_columnas_xlsx(blob, "A", "Z")

    def test_sugerir_columna(self):
        self.assertEqual(
            sugerir_columna(["Headline", "Resumen"], ("titulo", "headline")), "Headline")
        self.assertEqual(
            sugerir_columna(["Headline", "Resumen"], ("resumen", "resumen - aclaracion",
                                                     "cuerpoes")),
            "Resumen")
        self.assertIsNone(sugerir_columna(["A", "B"], ("titulo",)))

    def test_sugerir_columna_prioriza_exacta(self):
        # v4.29: "Subtitulo" no le gana a "Título" aunque venga antes.
        self.assertEqual(
            sugerir_columna(["Subtitulo", "Título"], ("título", "titulo", "headline")),
            "Título")
        self.assertEqual(
            sugerir_columna(["Titulo", "Resumen"], ("título", "titulo")), "Titulo")
        self.assertEqual(
            sugerir_columna(["TITULO", "RESUMEN"], ("resumen", "resumen - aclaracion")),
            "RESUMEN")

    def test_renombrado_pasa_por_el_loader_del_pipeline(self):
        # El xlsx renombrado debe ser legible por load_dossier_dataframe con
        # las columnas que el pipeline espera.
        from pipeline import load_dossier_dataframe
        blob = _xlsx_bytes(["Headline", "Body", "Medio"],
                           [["El título", "El cuerpo", "El Medio"]])
        out = renombrar_columnas_xlsx(blob, "Headline", "Body")
        df = load_dossier_dataframe(out)
        self.assertIn("Título", df.columns)
        self.assertIn("Resumen - Aclaracion", df.columns)
        self.assertEqual(df["Título"].iloc[0], "El título")
        self.assertEqual(df["Resumen - Aclaracion"].iloc[0], "El cuerpo")


if __name__ == "__main__":
    unittest.main()


class TestConfigCustom(unittest.TestCase):
    """v4.28: el ai_config de la pestaña nueva acepta PKL de tono/tema."""

    def _cfg(self, **kw):
        base = dict(brand="La Marca", alias_txt="Alias1; Alias2", voceros_txt="El Vocero",
                    criterio="Aspectual estricto", incluir_tema=False,
                    incluir_prominencia=False, enable_ai=True, api_key="k",
                    typesafe_api_key="tk", historial_dir=None,
                    tone_pkl_bytes=None, theme_pkl_bytes=None)
        base.update(kw)
        return construir_ai_config_custom(**base)

    def test_none_sin_analisis(self):
        self.assertIsNone(self._cfg(enable_ai=False))

    def test_con_pkl_tono_y_tema(self):
        cfg = self._cfg(tone_pkl_bytes=b"tono", theme_pkl_bytes=b"tema", incluir_tema=True)
        self.assertTrue(cfg["enabled"])
        self.assertEqual(cfg["brand"], "La Marca")
        self.assertEqual(cfg["aliases"], ["Alias1", "Alias2"])
        self.assertEqual(cfg["voceros"], ["El Vocero"])
        self.assertEqual(cfg["tone_pkl_bytes"], b"tono")
        self.assertEqual(cfg["theme_pkl_bytes"], b"tema")
        self.assertTrue(cfg["incluir_tema"])

    def test_solo_pkl_sin_ia_construye_config(self):
        cfg = self._cfg(enable_ai=False, tone_pkl_bytes=b"tono", api_key=None,
                        typesafe_api_key=None)
        self.assertIsNotNone(cfg)
        self.assertFalse(cfg["enabled"])
        self.assertIsNone(cfg["api_key"])
        self.assertEqual(cfg["tone_pkl_bytes"], b"tono")

    def test_solo_prominencia_sin_ia_ni_pkl(self):
        cfg = self._cfg(enable_ai=False, incluir_prominencia=True, api_key=None,
                        typesafe_api_key=None)
        self.assertIsNotNone(cfg)
        self.assertFalse(cfg["enabled"])
        self.assertTrue(cfg["incluir_prominencia"])
